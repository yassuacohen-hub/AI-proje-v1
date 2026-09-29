import openpyxl

wb = openpyxl.load_workbook('data/nace/sektor_meslek_nace_2026-05_resmi.xlsx')
ws = wb.active

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    nace_kodu_raw = str(row[4]).strip() if row[4] else ""
    if not nace_kodu_raw:
        continue
    nace_code = nace_kodu_raw.strip()
    if '47.79' in nace_code:
        print(f"NACE: {nace_code}")
        print(f"  Sektor: {row[0]} -> {row[1]}")
        print(f"  Meslek: {row[2]} -> {row[3]}")
        print(f"  NACE Tanim: {row[5]}")
        print()